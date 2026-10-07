"""Check that both poster builds quote exactly the generated numbers.

Every number the poster quotes comes from generated/poster_numbers.json and
carries three decimals (figure tick labels carry one or two). The check
passes only when:

* every three-decimal number in poster.pdf and poster.pptx is a generated
  value (a hand-typed number that drifted from the evidence fails here);
* both builds quote the same set of three-decimal numbers;
* the PDF is one 48 x 36 in page.

Run from the repository root:  python docs/poster/check_poster.py
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from pptx import Presentation
from pypdf import PdfReader

HERE = Path(__file__).resolve().parent
THREE_DECIMALS = re.compile(r"(?<![\d.])\d+\.\d{3}(?!\d)")


def _normalize(value: str) -> set[str]:
    return set(THREE_DECIMALS.findall(value.replace("−", "-")))


def pptx_text(path: Path) -> str:
    chunks = []
    for shape in Presentation(path).slides[0].shapes:
        if shape.has_text_frame:
            chunks.append(shape.text_frame.text)
        if getattr(shape, "has_table", False) and shape.has_table:
            chunks.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return "\n".join(chunks)


def main() -> int:
    numbers = json.loads((HERE / "generated" / "poster_numbers.json").read_text(encoding="utf-8"))["display"]
    allowed = set().union(*(_normalize(v) for v in numbers.values()))
    failures = []
    reader = PdfReader(HERE / "poster.pdf")
    page = reader.pages[0]
    size = (round(float(page.mediabox.width) / 72, 2), round(float(page.mediabox.height) / 72, 2))
    if len(reader.pages) != 1 or size != (48.0, 36.0):
        failures.append(f"poster.pdf must be one 48 x 36 in page, got {len(reader.pages)} page(s) of {size}")
    pdf_numbers = _normalize(page.extract_text())
    pptx_numbers = _normalize(pptx_text(HERE / "poster.pptx"))
    for name, found in (("poster.pdf", pdf_numbers), ("poster.pptx", pptx_numbers)):
        stray = found - allowed
        if stray:
            failures.append(f"{name} quotes numbers that are not generated: {sorted(stray)}")
    if pdf_numbers != pptx_numbers:
        failures.append(
            "the builds quote different numbers: only PDF "
            f"{sorted(pdf_numbers - pptx_numbers)}, only PPTX {sorted(pptx_numbers - pdf_numbers)}"
        )
    if failures:
        print("\n".join(failures))
        return 1
    print(f"ok: {len(pdf_numbers)} quoted numbers match across poster.pdf and poster.pptx; page {size[0]} x {size[1]} in")
    return 0


if __name__ == "__main__":
    sys.exit(main())
